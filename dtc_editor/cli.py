from __future__ import annotations
import argparse
import json
import os
from pathlib import Path
from datetime import datetime, timezone
from dtc_editor.pipeline import run_pipeline


def _run_holistic_mode(args):
    """Run the holistic (LLM-first) pipeline."""
    from dtc_editor.adapters.docx_adapter import extract_ir_and_inventory, emit_clean_docx, load_protected_terms
    from dtc_editor.holistic import (
        run_holistic_pipeline,
        HolisticConfig,
        generate_review_report,
    )

    # Load document
    print(f"Loading document: {args.input_docx}")
    ir, inventory = extract_ir_and_inventory(args.input_docx)

    # Load protected terms
    terms_path = Path(args.input_docx).parent / "protected_terms.txt"
    if terms_path.exists():
        protected_terms = load_protected_terms(str(terms_path))
    else:
        # Default protected terms for DTC documents
        protected_terms = {
            "Digital Twin Consortium", "DTC", "MEC", "ETSI", "IoT",
            "Internet of Things", "Kubernetes", "DePIN", "OT", "IT",
            "Multi-access Edge Computing", "Edge Computing", "5G",
        }

    # Find vale config
    vale_config = args.vale_config
    if not vale_config:
        default_vale = Path(__file__).parent.parent / ".vale.ini"
        if default_vale.exists():
            vale_config = str(default_vale)

    # Configure pipeline
    config = HolisticConfig(
        api_key=args.anthropic_api_key,
        model=args.llm_model,
        chunk_strategy=args.chunk_strategy,
        max_concurrent=2,  # Conservative for rate limits
        vale_config=vale_config,
        protected_terms=protected_terms,
        auto_accept=args.auto_accept,
    )

    # Run pipeline
    print(f"Running holistic pipeline (strategy: {args.chunk_strategy})")

    def progress(stage, completed, total):
        if total > 0:
            print(f"  {stage}: {completed}/{total}")

    result = run_holistic_pipeline(ir, config, progress_callback=progress)

    # Create output directory
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%SZ")
    doc_stem = Path(args.input_docx).stem
    bundle_dir = Path(args.out) / f"{doc_stem}_{timestamp}"
    bundle_dir.mkdir(parents=True, exist_ok=True)

    # Copy original
    import shutil
    orig_dest = bundle_dir / f"{doc_stem}.original.docx"
    shutil.copy2(args.input_docx, orig_dest)

    # Emit clean document
    clean_dest = bundle_dir / f"{doc_stem}.clean.docx"
    emit_clean_docx(args.input_docx, result.final_ir, str(clean_dest))

    # Optional: Add figure captions
    figure_stats = None
    if getattr(args, 'add_figure_captions', False):
        from dtc_editor.adapters.figure_captions import process_figure_captions, CaptionConfig
        caption_config = CaptionConfig(
            use_llm=True,
            api_key=args.anthropic_api_key,
        )
        figure_stats = process_figure_captions(
            str(clean_dest),  # Input: the clean docx we just created
            str(clean_dest),  # Output: overwrite same file
            caption_config,
        )
        print(f"Figure captions: {figure_stats['total_figures']} figures, {figure_stats['captions_inferred']} inferred, {figure_stats['placeholders_added']} placeholders")

    # Optional: Add TOC/TOF/TOT
    toc_stats = None
    if getattr(args, 'add_toc', False):
        from docx import Document
        from dtc_editor.adapters.document_restructure import (
            infer_document_structure,
            insert_inferred_toc,
            insert_inferred_tof,
            insert_inferred_tot,
        )

        doc = Document(str(clean_dest))
        inferred = infer_document_structure(doc)

        # Insert in reverse order (TOT, TOF, TOC) so positions stay valid
        insert_pos = 2  # After title

        tot_inserted = 0
        tof_inserted = 0
        toc_inserted = 0

        if inferred.table_entries:
            tot_inserted = insert_inferred_tot(doc, inferred.table_entries, insert_pos)
        if inferred.figure_entries:
            tof_inserted = insert_inferred_tof(doc, inferred.figure_entries, insert_pos)
        if inferred.toc_entries:
            toc_inserted = insert_inferred_toc(doc, inferred.toc_entries, insert_pos)

        doc.save(str(clean_dest))

        toc_stats = {
            "toc_entries": len(inferred.toc_entries),
            "tof_entries": len(inferred.figure_entries),
            "tot_entries": len(inferred.table_entries),
        }
        print(f"TOC: {toc_stats['toc_entries']} entries, TOF: {toc_stats['tof_entries']} figures, TOT: {toc_stats['tot_entries']} tables")

    # Generate review report
    review_report = generate_review_report(result)
    review_dest = bundle_dir / f"{doc_stem}.review.md"
    review_dest.write_text(review_report)

    # Also save to custom location if requested
    if args.review_file:
        Path(args.review_file).write_text(review_report)

    # Output summary
    output = {
        "mode": "holistic",
        "bundle_dir": str(bundle_dir),
        "chunk_strategy": args.chunk_strategy,
        "total_chunks": result.stats.total_chunks,
        "rewritable_chunks": result.stats.rewritable_chunks,
        "accepted": result.stats.accepted,
        "rejected": result.stats.rejected,
        "flagged": result.stats.flagged,
        "words_original": result.stats.total_words_original,
        "words_final": result.stats.total_words_final,
        "word_reduction": f"{(1 - result.stats.total_words_final / result.stats.total_words_original):.1%}",
        "processing_time_s": round(result.stats.total_time_s, 1),
        "review_needed": result.review_needed,
        "review_file": str(review_dest),
    }

    # Add figure stats if processed
    if figure_stats:
        output["figures_processed"] = figure_stats["total_figures"]
        output["captions_inferred"] = figure_stats["captions_inferred"]
        output["placeholders_added"] = figure_stats["placeholders_added"]

    # Add TOC stats if processed
    if toc_stats:
        output["toc_entries"] = toc_stats["toc_entries"]
        output["tof_entries"] = toc_stats["tof_entries"]
        output["tot_entries"] = toc_stats["tot_entries"]

    print(json.dumps(output, indent=2))


def main():
    ap = argparse.ArgumentParser(
        prog="dtc-edit",
        description="DTC Editorial Engine vNext2"
    )

    # Make input_docx optional to allow --gui without input
    ap.add_argument("input_docx", nargs="?", help="Path to input .docx")
    ap.add_argument("--out", default="./dtc_out", help="Output directory")
    ap.add_argument(
        "--mode", default="safe",
        choices=["safe", "rewrite", "holistic"],
        help="Run mode: safe (rules only), rewrite (rules + LLM fixes), holistic (LLM-first rewrite)"
    )
    ap.add_argument("--author", default="DTC Editorial Engine", help="Compare author name")
    ap.add_argument("--compare-backend", default=None, help="Prefer compare backend: aspose|word_com")

    # LLM options
    llm_group = ap.add_argument_group("LLM Options")
    llm_group.add_argument(
        "--use-llm",
        action="store_true",
        help="Enable LLM-based prose rewrites (requires --anthropic-api-key or ANTHROPIC_API_KEY env var)"
    )
    llm_group.add_argument(
        "--anthropic-api-key",
        default=os.environ.get("ANTHROPIC_API_KEY"),
        help="Anthropic API key (or set ANTHROPIC_API_KEY env var)"
    )
    llm_group.add_argument(
        "--llm-model",
        default="claude-sonnet-4-20250514",
        help="Claude model to use for rewrites (default: claude-sonnet-4-20250514)"
    )

    # Google Docs export options
    google_group = ap.add_argument_group("Google Docs Export")
    google_group.add_argument(
        "--export-google-docs",
        action="store_true",
        help="Upload clean.docx to Google Drive as Google Doc"
    )
    google_group.add_argument(
        "--google-credentials",
        help="Path to Google service account or OAuth credentials JSON"
    )
    google_group.add_argument(
        "--google-folder-id",
        help="Target Google Drive folder ID (optional)"
    )

    # Vale options
    vale_group = ap.add_argument_group("Vale Linting")
    vale_group.add_argument(
        "--use-vale",
        action="store_true",
        help="Enable Vale prose linting"
    )
    vale_group.add_argument(
        "--vale-config",
        help="Path to .vale.ini config file (optional)"
    )

    # Holistic mode options
    holistic_group = ap.add_argument_group("Holistic Mode Options (--mode holistic)")
    holistic_group.add_argument(
        "--chunk-strategy",
        default="paragraph",
        choices=["paragraph", "section", "adaptive"],
        help="How to chunk document: paragraph (safest), section (coherent), adaptive (balanced)"
    )
    holistic_group.add_argument(
        "--auto-accept",
        action="store_true",
        help="Auto-accept all passing validations without human review"
    )
    holistic_group.add_argument(
        "--review-file",
        help="Output path for review report (markdown)"
    )
    holistic_group.add_argument(
        "--add-figure-captions",
        action="store_true",
        help="Detect figures and add missing captions (uses LLM if available)"
    )
    holistic_group.add_argument(
        "--add-toc",
        action="store_true",
        help="Add Table of Contents, Figures, and Tables"
    )

    # GUI option
    ap.add_argument(
        "--gui",
        action="store_true",
        help="Launch graphical user interface"
    )

    args = ap.parse_args()

    # Launch GUI if requested
    if args.gui:
        from dtc_editor.gui import launch_gui
        launch_gui()
        return

    # Require input_docx for CLI mode
    if not args.input_docx:
        ap.error("input_docx is required (or use --gui for graphical interface)")

    # Validate LLM options
    if args.use_llm and not args.anthropic_api_key:
        ap.error("--use-llm requires --anthropic-api-key or ANTHROPIC_API_KEY environment variable")

    # Holistic mode requires API key
    if args.mode == "holistic" and not args.anthropic_api_key:
        ap.error("--mode holistic requires --anthropic-api-key or ANTHROPIC_API_KEY environment variable")

    # Validate Google options
    if args.export_google_docs and not args.google_credentials:
        ap.error("--export-google-docs requires --google-credentials")

    # Run holistic pipeline if requested
    if args.mode == "holistic":
        _run_holistic_mode(args)
        return

    payload = run_pipeline(
        input_docx=args.input_docx,
        out_dir=args.out,
        mode=args.mode,
        author=args.author,
        prefer_compare_backend=args.compare_backend,
        # LLM options
        use_llm=args.use_llm,
        anthropic_api_key=args.anthropic_api_key,
        llm_model=args.llm_model,
        # Google options
        export_google_docs=args.export_google_docs,
        google_credentials=args.google_credentials,
        google_folder_id=args.google_folder_id,
        # Vale options
        use_vale=args.use_vale,
        vale_config_path=args.vale_config,
    )

    # Build output summary
    output = {
        "bundle_dir": payload["artifacts"]["clean_docx"].rsplit("/", 1)[0],
        "redline_status": payload["redline_engine"]["status"],
        "editops_total": payload["stats"]["editops_total"],
        "editops_applied": payload["stats"]["editops_applied"],
        "editops_llm": payload["stats"].get("editops_llm", 0),
        "editops_vale": payload["stats"].get("editops_vale", 0),
        "findings_total": payload["stats"]["findings_total"],
        "persnickety_ok": payload["persnicketybot"]["ok"],
    }

    # Add Vale status if used
    if payload.get("vale"):
        output["vale_status"] = payload["vale"]["status"]

    # Add Google Docs URL if available
    if payload.get("google_export") and payload["google_export"].get("status") == "ok":
        output["google_doc_url"] = payload["google_export"]["web_view_link"]

    print(json.dumps(output, indent=2))


if __name__ == "__main__":
    main()
